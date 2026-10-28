"""Builds presentation/slides.pptx from results/ (numbers are read, not typed)."""
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

R = "results/"
m = pd.read_csv(R + "metrics.csv").set_index("model")
tr = pd.read_csv("data/processed/train.csv")
nitems = len(pd.read_csv("data/processed/items.csv"))
NAVY, ORANGE, GREY = RGBColor(0x1F, 0x2A, 0x44), RGBColor(0xE8, 0x74, 0x3B), RGBColor(0x55, 0x5F, 0x6E)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def text(sl, x, y, w, h, s, size=18, bold=False, color=GREY):
    tf = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for n, line in enumerate(s if isinstance(s, list) else [s]):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.text, p.space_after = line, Pt(10)
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
    return tf


def slide(title, notes=""):
    sl = prs.slides.add_slide(blank)
    bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text(sl, 0.6, 0.2, 12, 0.8, title, 30, True, RGBColor(255, 255, 255))
    if notes: sl.notes_slide.notes_text_frame.text = notes
    return sl


def bullets(sl, items, x=0.7, y=1.5, w=11.9, size=22):
    text(sl, x, y, w, 5, ["•  " + i for i in items], size)


# 1 title
sl = prs.slides.add_slide(blank)
bg = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(sl, 0.8, 2.2, 11.5, 1.6, "Retail Product Recommendation Using GNN", 44, True, RGBColor(255, 255, 255))
text(sl, 0.8, 3.7, 11.5, 0.8, "LightGCN on Amazon Reviews 2023 (Video Games)", 24, False, ORANGE)
text(sl, 0.8, 5.3, 11.5, 1.2, ["[Team member names]  |  [Enrollment numbers]", "Guide: [Guide name]  |  [Institute / Department / Year]"], 18, False, RGBColor(0xCF, 0xD6, 0xE2))

# 2 problem
sl = slide("Problem: what should this customer see next?")
bullets(sl, ["Retail catalogues hold thousands of products; customers see a handful.",
             "Goal: given a customer ID (or a few liked products) return a ranked list of products they are likely to want.",
             "Signal: who bought what, and who buys like whom - naturally a graph.",
             "Output: top-K candidates, history excluded, judged on future behaviour."])

# 3 data
sl = slide("Data: Amazon Reviews 2023 - Video Games")
bullets(sl, ["Official 5-core files, already split by timestamp: train < valid < test.",
             "Rating >= 4 = positive implicit interaction.",
             f"Subsample for laptop CPU: {tr.user.nunique():,} customers, {nitems:,} products, {len(tr):,} training edges.",
             "Users with future activity kept first, so the test set is meaningful (graph still uses train edges only).",
             "Metadata (title, category, price) used for display in the demo."])

# 4 graph
sl = slide("Graph: customers and products as nodes")
bullets(sl, ["Bipartite graph; edge = positive review in the TRAINING period.",
             "A = [[0, R], [R^T, 0]]  →  A_hat = D^-1/2 · A · D^-1/2  (sparse).",
             f"Density {len(tr)/(tr.user.nunique()*nitems):.4f}; ~{len(tr)/tr.user.nunique():.1f} edges per customer, long-tailed item degrees.",
             "No validation/test edge ever enters the graph → no leakage."])

# 5 model
sl = slide("Model: LightGCN")
bullets(sl, ["Propagate:  E(k+1) = A_hat · E(k)   - no weights, no non-linearity.",
             "Final embedding = mean of layers 0..K  (K = 3, dim 64).",
             "Score = dot(user, item).",
             "Loss (BPR):  -log σ(s(u,i) - s(u,j))  + L2 on initial embeddings; Adam, early stopping on val NDCG@10.",
             "Baselines: Popularity, and MF-BPR (same loss, no graph)."])

# 6 protocol
sl = slide("Evaluation protocol")
bullets(sl, ["Full ranking: score ALL products for each test customer.",
             "Mask products already seen in train/val; take the top K = 5, 10, 20.",
             "Precision@K, Recall@K, NDCG@K averaged over customers with test positives.",
             f"{int(m.loc['lightgcn','n_eval_users']):,} test customers; metrics unit-tested on hand-computed cases."])

# 7 results table
sl = slide("Results (test set)")
rows = [("Metric", "Popularity", "MF-BPR", "LightGCN")] + [
    (c.upper(),
     *[f"{m.loc[k, c]:.4f}" for k in ("popularity", "mf_bpr", "lightgcn")])
    for c in ("ndcg@10", "recall@10", "precision@10", "ndcg@20", "recall@20")]
tb = sl.shapes.add_table(len(rows), 4, Inches(1.2), Inches(1.6), Inches(10.9), Inches(3.4)).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tb.cell(i, j); c.text = v
        c.text_frame.paragraphs[0].font.size = Pt(20)
        c.text_frame.paragraphs[0].font.bold = (j == 3 or i == 0)
g, p = m.loc["lightgcn", "ndcg@10"], m.loc["popularity", "ndcg@10"]
text(sl, 1.2, 5.4, 11, 1.5, f"LightGCN is best on every metric: NDCG@10 is {100*(g/p-1):.0f}% above popularity. "
     "Absolute scores are low (sparse data, future period); one seed, so small gaps are indicative only.", 18)

# 8 plots
sl = slide("Comparison and training behaviour")
sl.shapes.add_picture(R + "model_comparison.png", Inches(0.4), Inches(1.4), width=Inches(7.6))
sl.shapes.add_picture(R + "training_curves.png", Inches(8.1), Inches(1.5), width=Inches(5.0))
text(sl, 0.6, 4.6, 12, 2.5, ["MF-BPR overfits: loss falls, validation NDCG stalls - too little data per customer for a free vector.",
                              "LightGCN's neighbourhood averaging smooths sparse customers and generalises better (our interpretation; no ablation run)."], 18)

# 9 demo
sl = slide("Live demo (Streamlit)")
sl.shapes.add_picture(R + "app_tab1.png", Inches(0.5), Inches(1.3), width=Inches(6.2))
sl.shapes.add_picture(R + "app_tab2.png", Inches(6.9), Inches(1.3), width=Inches(6.2))
text(sl, 0.6, 5.9, 12, 1.2, "Left: customer ID → past purchases next to recommendations.   Right: cold start - pick products, average their embeddings, rank the rest.", 16)

# 10 limits
sl = slide("Limitations and future scope")
bullets(sl, ["One category, subsampled; single seed, no confidence intervals.",
             "Only interaction edges used - no text, category or price features yet.",
             "Cold-start works for new customers, not for brand-new products.",
             "Next: side-information features, hard negatives, multi-seed evaluation, other graph models, full-data run."], size=22)

# 11 conclusion
sl = slide("Conclusion")
bullets(sl, ["End-to-end graph recommender: data → graph → LightGCN → evaluation → app.",
             "Leakage-free temporal protocol with full-ranking metrics.",
             f"LightGCN beat Popularity and MF-BPR (NDCG@10 {g:.4f} vs {p:.4f} / {m.loc['mf_bpr','ndcg@10']:.4f}).",
             "Thank you - questions?"], size=24)
prs.save("presentation/slides.pptx")
print("saved", len(prs.slides._sldIdLst), "slides")
